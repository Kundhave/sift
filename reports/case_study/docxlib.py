"""Small python-docx layer: IEEE-Access-like page (two columns, Times New Roman 10 pt), numbered headings,
SEQ-field captions, real tables, equations, and full-width blocks via continuous section breaks."""
from __future__ import annotations
import re
from contextlib import contextmanager

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
COL_W = 3.4
PAGE_W = 7.0
ITALIC_EXCEPT = {"log", "exp", "max", "min", "arg", "argmax", "KL", "NDCG", "ECE", "CE", "IG", "softmax", "cos", "sup", "std", "mean", "top", "rand", "if", "else", "for", "and", "with", "where", "in", "of", "to", "d", "G", "D"}
_TAG = re.compile(r"(</?(?:b|i|sub|sup|u|code)>)")


def _set_font(run, size=None, bold=None, italic=None, color=None, name=FONT, small_caps=None, mono=False):
    run.font.name = "Courier New" if mono else name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), "Courier New" if mono else name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if small_caps is not None: run.font.small_caps = small_caps
    if color: run.font.color.rgb = RGBColor.from_string(color)


def _fmt(par, align=None, before=0, after=0, first=None, left=None, line=None, keep_next=False, keep_lines=False):
    pf = par.paragraph_format
    if align is not None: par.alignment = align
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if first is not None: pf.first_line_indent = Inches(first)
    if left is not None: pf.left_indent = Inches(left)
    if line: pf.line_spacing = line
    pf.keep_with_next = keep_next
    pf.keep_together = keep_lines
    pf.widow_control = True


TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]
TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden"]


def _reorder(el, order):
    def key(ch):
        tag = ch.tag.split("}")[1]
        return order.index(tag) if tag in order else len(order)
    kids = sorted(list(el), key=key)
    for k in list(el): el.remove(k)
    for k in kids: el.append(k)


def fix_table_xml(tbl):
    _reorder(tbl.tblPr, TBLPR_ORDER)
    for tc in tbl.iter(qn("w:tcPr")): _reorder(tc, TCPR_ORDER)
    for tr in tbl.iter(qn("w:trPr")): _reorder(tr, TRPR_ORDER)


class Doc:
    def __init__(self):
        self.d = Document()
        d = self.d
        s = d.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.left_margin = s.right_margin = Inches(0.75)
        s.top_margin, s.bottom_margin = Inches(0.8), Inches(0.9)
        s.header_distance, s.footer_distance = Inches(0.4), Inches(0.45)
        st = d.styles["Normal"]
        st.font.name = FONT; st.font.size = Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(0)
        for name, size, bold, ital, caps in (("Heading 1", 10, True, False, True), ("Heading 2", 10, True, True, False), ("Heading 3", 10, False, True, False)):
            h = d.styles[name]
            h.font.name = FONT; h.font.size = Pt(size); h.font.bold = bold; h.font.italic = ital; h.font.small_caps = caps
            h.font.color.rgb = RGBColor(0, 0, 0)
            rpr = h.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(a), FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
        self.fig_n = 0; self.tab_n = 0; self.h1 = 0; self.h2 = 0; self.eq_n = 0
        self.body_words = {}
        self.cur = None
        self._wide_depth = 0
        self.in_wide = False
        self.wide_count = 0

    # ---------------------------------------------------------------- sections
    def _cols(self, sec, n, space=0.2):
        sp = sec._sectPr
        c = sp.find(qn("w:cols"))
        if c is None:
            c = OxmlElement("w:cols"); sp.append(c)
        c.set(qn("w:num"), str(n)); c.set(qn("w:space"), str(int(space * 1440)))
        if n == 1 and c.get(qn("w:equalWidth")) is not None: del c.attrib[qn("w:equalWidth")]

    def start_two_columns(self):
        sec = self.d.add_section(WD_SECTION.CONTINUOUS)
        self._cols(sec, 2)

    @contextmanager
    def wide(self):
        """Full-width block inside the two-column body."""
        sec = self.d.add_section(WD_SECTION.CONTINUOUS); self._cols(sec, 1)
        self.wide_count += 1
        self.in_wide = True
        try:
            yield
        finally:
            self.in_wide = False
            sec2 = self.d.add_section(WD_SECTION.CONTINUOUS); self._cols(sec2, 2)

    # ---------------------------------------------------------------- inline text
    def runs(self, par, text, size=10, bold=False, italic=False, color=None, mono=False):
        b, i, sub, sup, code = bold, italic, False, False, False
        for tok in _TAG.split(text):
            if not tok: continue
            if tok == "<b>": b = True
            elif tok == "</b>": b = bold
            elif tok == "<i>": i = True
            elif tok == "</i>": i = italic
            elif tok == "<sub>": sub = True
            elif tok == "</sub>": sub = False
            elif tok == "<sup>": sup = True
            elif tok == "</sup>": sup = False
            elif tok == "<code>": code = True
            elif tok == "</code>": code = False
            else:
                r = par.add_run(tok)
                _set_font(r, size=size - 1.2 if code else size, bold=b, italic=i, color=color, mono=code)
                if sub: r.font.subscript = True
                if sup: r.font.superscript = True
        return par

    def p(self, text, size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=0.18, before=0, after=2.5, bold=False, italic=False, keep_next=False, left=None, count_key=None):
        par = self.d.add_paragraph()
        _fmt(par, align, before, after, first, left=left, keep_next=keep_next)
        self.runs(par, text, size=size, bold=bold, italic=italic)
        if count_key is None: count_key = self.cur
        if count_key: self.body_words[count_key] = self.body_words.get(count_key, 0) + len(re.sub(r"<[^>]+>", "", text).split())
        return par

    def bullets(self, items, size=10, marker="•", after=1.5):
        for it in items:
            par = self.d.add_paragraph()
            _fmt(par, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, after, first=-0.17, left=0.2)
            self.runs(par, f"{marker}\t" + it, size=size)
            par.paragraph_format.tab_stops.add_tab_stop(Inches(0.2))
            if self.cur: self.body_words[self.cur] = self.body_words.get(self.cur, 0) + len(re.sub(r"<[^>]+>", "", it).split())

    def numbered(self, items, size=10, fmt="({})", after=1.5):
        for n, it in enumerate(items, 1):
            par = self.d.add_paragraph()
            _fmt(par, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, after, first=-0.25, left=0.28)
            self.runs(par, fmt.format(n) + "\t" + it, size=size)
            par.paragraph_format.tab_stops.add_tab_stop(Inches(0.28))
            if self.cur: self.body_words[self.cur] = self.body_words.get(self.cur, 0) + len(re.sub(r"<[^>]+>", "", it).split())

    # ---------------------------------------------------------------- headings
    def heading1(self, text, numbered=True, count_key=None):
        if numbered:
            self.h1 += 1; self.h2 = 0; text = f"{self.h1}. {text}"
        par = self.d.add_paragraph(style="Heading 1")
        _fmt(par, WD_ALIGN_PARAGRAPH.LEFT, 9, 3, keep_next=True)
        r = par.add_run(text); _set_font(r, 10, True, False, small_caps=True)
        self.cur = count_key or text
        return par

    def heading2(self, text, numbered=True, count_key=None):
        if numbered:
            self.h2 += 1; text = f"{self.h1}.{self.h2} {text}"
        par = self.d.add_paragraph(style="Heading 2")
        _fmt(par, WD_ALIGN_PARAGRAPH.LEFT, 5, 2, keep_next=True)
        r = par.add_run(text); _set_font(r, 10, True, True)
        if count_key: self.cur = count_key
        return par

    def run_in(self, head, text, **kw):
        """Paragraph that starts with a bold-italic run-in title."""
        return self.p(f"<b><i>{head}.</i></b> " + text, **kw)

    # ---------------------------------------------------------------- fields and captions
    def _seq(self, par, name, n, size=8, bold=True):
        def fc(t):
            r = par.add_run(); _set_font(r, size, bold)
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r._element.append(e)
        fc("begin")
        r = par.add_run(); _set_font(r, size, bold)
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" SEQ {name} \\* ARABIC "; r._element.append(it)
        fc("separate")
        r = par.add_run(str(n)); _set_font(r, size, bold)
        fc("end")

    def caption_table(self, text):
        self.tab_n += 1
        par = self.d.add_paragraph()
        _fmt(par, WD_ALIGN_PARAGRAPH.CENTER, 5, 2.5, keep_next=True, keep_lines=True)
        r = par.add_run("TABLE "); _set_font(r, 8, True)
        self._seq(par, "Table", self.tab_n)
        r = par.add_run(". "); _set_font(r, 8, True)
        self.runs(par, text, size=8, bold=False)
        return self.tab_n

    def caption_figure(self, text, container=None):
        self.fig_n += 1
        par = (container or self.d).add_paragraph()
        _fmt(par, WD_ALIGN_PARAGRAPH.JUSTIFY if len(text) > 110 else WD_ALIGN_PARAGRAPH.CENTER, 2, 6, keep_lines=True)
        r = par.add_run("FIGURE "); _set_font(r, 8, True)
        self._seq(par, "Figure", self.fig_n)
        r = par.add_run(". "); _set_font(r, 8, True)
        self.runs(par, text, size=8)
        return self.fig_n

    def picture(self, path, width):
        par = self.d.add_paragraph()
        _fmt(par, WD_ALIGN_PARAGRAPH.CENTER, 3, 0, keep_next=True)
        par.add_run().add_picture(str(path), width=Inches(width))

    def boxed_figure(self, path, width, caption):
        """Picture and caption inside one non-splitting borderless table cell, so they cannot be separated across columns."""
        t = self.d.add_table(rows=1, cols=1)
        t.autofit = False
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr
        el = OxmlElement("w:tblLayout"); el.set(qn("w:type"), "fixed"); tblPr.append(el)
        tw = tblPr.find(qn("w:tblW"))
        if tw is None:
            tw = OxmlElement("w:tblW"); tblPr.append(tw)
        tw.set(qn("w:w"), str(int(COL_W * 1440))); tw.set(qn("w:type"), "dxa")
        for gc in t._tbl.tblGrid.findall(qn("w:gridCol")): gc.set(qn("w:w"), str(int(COL_W * 1440)))
        mar = OxmlElement("w:tblCellMar")
        for side in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), "0"); e.set(qn("w:type"), "dxa"); mar.append(e)
        tblPr.append(mar)
        tr = t.rows[0]; tr._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cell = tr.cells[0]; cell.width = Inches(COL_W)
        fix_table_xml(t._tbl)
        par = cell.paragraphs[0]
        _fmt(par, WD_ALIGN_PARAGRAPH.CENTER, 3, 0, keep_next=True)
        par.add_run().add_picture(str(path), width=Inches(min(width, COL_W - 0.2)))
        n = self.caption_figure(caption, container=cell)
        sp = self.d.add_paragraph(); _fmt(sp, None, 0, 3); sp.paragraph_format.line_spacing = Pt(3)
        return n

    def figure(self, path, width, caption):
        self.picture(path, width)
        return self.caption_figure(caption)

    # ---------------------------------------------------------------- tables
    def table(self, rows, widths, caption=None, size=7, aligns=None, bold_rows=(), bold_first_col=False, shade_rows=()):
        if caption: self.caption_table(caption)
        maxw = PAGE_W if self.in_wide else COL_W
        if sum(widths) > maxw:
            f = maxw / sum(widths); widths = [w * f for w in widths]
        ncol = len(widths)
        t = self.d.add_table(rows=len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        for tag in ("w:tblLayout",):
            el = OxmlElement(tag); el.set(qn("w:type"), "fixed"); tblPr.append(el)
        bd = OxmlElement("w:tblBorders")
        for side, sz, col in (("top", 8, "000000"), ("bottom", 8, "000000"), ("insideH", 2, "BBBBBB")):
            e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), col); bd.append(e)
        tblPr.append(bd)
        mar = OxmlElement("w:tblCellMar")
        for side, w in (("top", 15), ("left", 45), ("bottom", 15), ("right", 45)):
            e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(w)); e.set(qn("w:type"), "dxa"); mar.append(e)
        tblPr.append(mar)
        tw = tblPr.find(qn("w:tblW"))
        if tw is None:
            tw = OxmlElement("w:tblW"); tblPr.append(tw)
        tw.set(qn("w:w"), str(int(sum(widths) * 1440))); tw.set(qn("w:type"), "dxa")
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(int(w * 1440)))
        aligns = aligns or ["l"] + ["c"] * (ncol - 1)
        amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT, "j": WD_ALIGN_PARAGRAPH.JUSTIFY}
        for ri, row in enumerate(rows):
            tr = t.rows[ri]
            trPr = tr._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit"); trPr.append(cs)
            if ri == 0:
                th = OxmlElement("w:tblHeader"); trPr.append(th)
            for ci, val in enumerate(row):
                cell = tr.cells[ci]
                cell.width = Inches(widths[ci])
                par = cell.paragraphs[0]
                _fmt(par, amap[aligns[ci]] if ri else (WD_ALIGN_PARAGRAPH.LEFT if aligns[ci] == "l" else WD_ALIGN_PARAGRAPH.CENTER), 0, 0, keep_lines=True)
                if ri < len(rows) - 1: par.paragraph_format.keep_with_next = True
                b = ri == 0 or ri in bold_rows or (bold_first_col and ci == 0)
                self.runs(par, str(val), size=size, bold=b)
                if ri == 0 or ri in shade_rows:
                    tcPr = cell._tc.get_or_add_tcPr()
                    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), "E7ECF3" if ri == 0 else "F3F3F3"); tcPr.append(sh)
                if ri == 0:
                    tcPr = cell._tc.get_or_add_tcPr()
                    bb = OxmlElement("w:tcBorders"); e = OxmlElement("w:bottom"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "000000"); bb.append(e); tcPr.append(bb)
        fix_table_xml(t._tbl)
        sp = self.d.add_paragraph(); _fmt(sp, None, 0, 4)
        sp.paragraph_format.line_spacing = Pt(4)
        return t

    # ---------------------------------------------------------------- equations
    def equation(self, text, wide=False):
        """text uses <i>, <sub>, <sup>; the number is added at the right margin. Variables should already be wrapped in <i>."""
        self.eq_n += 1
        par = self.d.add_paragraph()
        width = PAGE_W if wide else COL_W
        _fmt(par, WD_ALIGN_PARAGRAPH.LEFT, 3, 4, keep_lines=True)
        ts = par.paragraph_format.tab_stops
        from docx.enum.text import WD_TAB_ALIGNMENT
        ts.add_tab_stop(Inches(width / 2 - 0.0), WD_TAB_ALIGNMENT.CENTER); ts.add_tab_stop(Inches(width), WD_TAB_ALIGNMENT.RIGHT)
        r = par.add_run("\t"); _set_font(r, 10)
        self.runs(par, text, size=10)
        r = par.add_run(f"\t({self.eq_n})"); _set_font(r, 10)
        return self.eq_n

    # ---------------------------------------------------------------- header/footer
    def header_footer(self, running_title):
        sec = self.d.sections[0]
        hp = sec.header.paragraphs[0]; _fmt(hp, WD_ALIGN_PARAGRAPH.CENTER, 0, 0)
        r = hp.add_run(running_title); _set_font(r, 8, False, True, color="555555")
        fp = sec.footer.paragraphs[0]; _fmt(fp, WD_ALIGN_PARAGRAPH.CENTER, 0, 0)
        for t, txt in (("begin", None), ("instr", " PAGE "), ("separate", None), ("text", "1"), ("end", None)):
            r = fp.add_run(); _set_font(r, 8.5)
            if t == "instr":
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._element.append(it)
            elif t == "text":
                r.text = txt
            else:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r._element.append(e)

    def save(self, path):
        self.d.save(str(path))
